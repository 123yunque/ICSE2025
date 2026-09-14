from pathlib import Path
import re, json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE

ROOT = Path(__file__).resolve().parent
NAME = '代码大模型文献读后报告_2025-2026'
content = (ROOT / (NAME + '.md')).read_text(encoding='utf-8')
doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = section.bottom_margin = Inches(0.8)
section.left_margin = section.right_margin = Inches(0.85)
section.header_distance = section.footer_distance = Inches(0.35)

def font(style, size, bold=False):
    style.font.name = 'Calibri'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0,0,0)
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '宋体')

font(doc.styles['Normal'], 11)
normal = doc.styles['Normal'].paragraph_format
normal.line_spacing = 1.2
normal.space_after = Pt(7)
normal.widow_control = True
for name,size in [('Title',22),('Subtitle',15),('Heading 1',16),('Heading 2',13),('Heading 3',11.5)]:
    font(doc.styles[name],size,name != 'Subtitle')
    s=doc.styles[name].paragraph_format
    s.keep_with_next=True
    s.space_before=Pt(10)
    s.space_after=Pt(8)
# Keep title treatments plain and academic; remove Word's decorative title border.
title_ppr = doc.styles['Title'].element.get_or_add_pPr()
title_border = title_ppr.find(qn('w:pBdr'))
if title_border is not None:
    title_ppr.remove(title_border)
font(doc.styles['Caption'],9)
section.different_first_page_header_footer = True
hp=section.header.paragraphs[0]
hp.text='代码大模型基准可信性与执行推理评测｜文献读后报告'
hp.style='Caption'
fp=section.footer.paragraphs[0]
fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=fp.add_run('— ')
r.font.size=Pt(9)
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE')
fp._p.append(field)
fp.add_run(' —').font.size=Pt(9)

def inline(p,text):
    pos=0
    for m in re.finditer(r'\[([^\]]+)\]\((https?://[^)]+)\)',text):
        if m.start()>pos: p.add_run(text[pos:m.start()])
        h=OxmlElement('w:hyperlink')
        rid=p.part.relate_to(m.group(2),RELATIONSHIP_TYPE.HYPERLINK,is_external=True)
        h.set(qn('r:id'),rid)
        rr=OxmlElement('w:r'); props=OxmlElement('w:rPr')
        co=OxmlElement('w:color');co.set(qn('w:val'),'174A70');props.append(co)
        u=OxmlElement('w:u');u.set(qn('w:val'),'single');props.append(u)
        rr.append(props);tt=OxmlElement('w:t');tt.text=m.group(1);rr.append(tt);h.append(rr);p._p.append(h)
        pos=m.end()
    if pos<len(text):p.add_run(text[pos:])

def table(rows):
    t=doc.add_table(rows=1,cols=len(rows[0]))
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.style='Table Grid'
    for i,row in enumerate(rows):
        cells=t.rows[0].cells if i==0 else t.add_row().cells
        for c,v in zip(cells,row):
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=c.paragraphs[0];p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(4)
            p.paragraph_format.line_spacing=1.1
            inline(p,v.strip())
            for r in p.runs:r.font.size=Pt(9);r.bold=i==0
            if i==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EDEDED');c._tc.get_or_add_tcPr().append(sh)
        trPr=t.rows[i]._tr.get_or_add_trPr()
        cant=OxmlElement('w:cantSplit');trPr.append(cant)
        if i==0:
            repeat=OxmlElement('w:tblHeader');trPr.append(repeat)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

lines=content.splitlines(); i=0; heading_count=0
while i<len(lines):
    line=lines[i].strip();i+=1
    if not line:continue
    if line.startswith('|'):
        rows=[[v.strip() for v in line.strip('|').split('|')]]
        while i<len(lines) and lines[i].strip().startswith('|'):
            ll=lines[i].strip();i+=1
            if re.fullmatch(r'[| :\-]+',ll):continue
            rows.append([v.strip() for v in ll.strip('|').split('|')])
        table(rows);continue
    if line.startswith('# '):
        p=doc.add_paragraph(line[2:],'Title');p.paragraph_format.space_before=Pt(64)
    elif line.startswith('## '):
        title=line[3:]
        if title.startswith('2025'):
            p=doc.add_paragraph(title,'Subtitle');p.paragraph_format.space_after=Pt(38)
        else:
            p=doc.add_paragraph(title,'Heading 1')
            # Section 4 follows a short concluding paragraph; flowing the heading
            # onto that page prevents an isolated paragraph page.
            p.paragraph_format.page_break_before = title != '四、对开题方案的具体建议'
    elif line.startswith('### '):
        title=line[4:];p=doc.add_paragraph(title,'Heading 2')
        if re.match(r'\[(?:[2-9]|1[0-9]|20)\]',title):p.paragraph_format.page_break_before=True
        heading_count+=1
    else:
        p=doc.add_paragraph();inline(p,line)

doc.core_properties.title='代码大模型基准可信性与执行推理评测：2025—2026 年国际文献读后报告'
doc.core_properties.subject='20 篇论文研读、版本核查、方法比较与开题研究建议'
doc.core_properties.author=''
doc.core_properties.keywords='代码大模型,基准污染,执行推理,多粒度评测'
out=ROOT/(NAME+'.docx');doc.save(out)
info={'docx':str(out),'markdown':str(ROOT/(NAME+'.md')),'characters':len(content),'paper_reviews':len(re.findall(r'^### \[\d+\]',content,re.M)),'tables':len(doc.tables),'links':len(re.findall(r'\]\(https?://',content)),'local_pdf_count':len(list((ROOT/'sources').glob('*.pdf'))),'evidence_A':17,'evidence_B':[3,9,12]}
(ROOT/'report_check.json').write_text(json.dumps(info,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(info,ensure_ascii=True))
